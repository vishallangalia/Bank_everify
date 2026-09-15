"""
Income Tax e-Filing Portal — Add Bank Account + E-Verify Automation
=====================================================================

FLOW (jaisa recording + screenshots mein dekha):
  1. Script login page kholta hai. USER MANUALLY PAN/User-ID daalega aur
     "Continue" dabayega (yeh step automate NAHI hai).
  2. Password page detect hote hi:
       - Password page par dikhne wala PAN read karke Excel (Bank.xlsx)
         mein uska password dhoondta hai.
       - "Please confirm your secure access message displayed above" checkbox
         tick karta hai.
       - Password field fill karta hai.
       - "Continue" click karta hai.
     Agar "Dual Login Detected" popup aaye (session kahin aur active hai), to
     "Login Here" button khud dabata hai.
  3. Dashboard load hone ke baad:
       Profile dropdown (upar right, naam ke neeche) -> "My Profile"
  4. Personal Details page par -> left sidebar mein "My Bank Account"
  5. "My Bank Accounts" page par -> "+ Add Bank Account"
  6. Add Bank Account form:
       - Bank Account Number + Confirm Bank Account Number (Excel se)
       - Account Type = "Saving Bank Account" (hamesha)
       - Account Holder Type = "Primary" (hamesha)
       - IFSC (Excel se) -> Bank Name/Branch auto-populate hone ka wait
       - "Nominate this bank account for Refund" = "Yes"
       - "Proceed To E-Verify" click
  7. e-Verify page par "OTP on mobile number registered with Aadhaar" radio
     select karke "Continue" click karta hai.
  8. Popup modal mein "I agree to validate my Aadhaar Details" checkbox tick
     karke "Generate Aadhaar OTP" click karta hai.
  9. "Verify OTP" modal aa jayega — YAHAN SE USER MANUALLY OTP daalega aur
     "Validate" dabayega. Script yahin ruk jata hai, browser open rehta hai.

Requirements:
    pip install playwright openpyxl
    playwright install chromium

Usage:
    python main.py

Debugging: agar koi step fail ho jaye, script apne aap "debug/" folder mein
screenshot + page HTML save kar deta hai (error_<timestamp>.png/.html).
"""

import re
import sys
import time
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright, Page, TimeoutError as PWTimeoutError

# ----------------------------- CONFIG ---------------------------------
LOGIN_URL = "https://eportal.incometax.gov.in/iec/foservices/#/login"
EXCEL_PATH = Path(__file__).parent / "Bank.xlsx"
NAV_TIMEOUT_MS = 5 * 60 * 1000   # manual steps ke liye generous wait (5 min)
ACTION_TIMEOUT_MS = 30 * 1000    # normal automated clicks ke liye
ACCOUNT_TYPE_LABEL = "Saving Bank Account"     # form mein hamesha yahi select hoga
ACCOUNT_HOLDER_TYPE = "Primary"                # form mein hamesha yahi select hoga
# ------------------------------------------------------------------------


def load_partner_map(excel_path: Path) -> dict:
    """Excel se PAN -> {password, bank_account, ifsc} mapping load karta hai
    (case-insensitive PAN). Columns expected: PAN, Password, Bank Account No,
    IFSC Code (header naam case-insensitive match hota hai)."""
    if not excel_path.exists():
        print(f"[ERROR] Excel file nahi mili: {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    header = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    pan_col = col("pan")
    pwd_col = col("password")
    bank_col = col("bank account no", "bank account number")
    ifsc_col = col("ifsc code", "ifsc")

    missing = [n for n, c in [("PAN", pan_col), ("Password", pwd_col),
                               ("Bank Account No", bank_col), ("IFSC Code", ifsc_col)] if c is None]
    if missing:
        print(f"[ERROR] Excel mein ye column(s) nahi mile: {', '.join(missing)}")
        sys.exit(1)

    mapping = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[pan_col] is None:
            continue
        pan = str(row[pan_col]).strip().upper()
        if not pan:
            continue
        mapping[pan] = {
            "password": str(row[pwd_col]).strip() if row[pwd_col] is not None else "",
            "bank_account": str(row[bank_col]).strip() if row[bank_col] is not None else "",
            "ifsc": str(row[ifsc_col]).strip().upper() if row[ifsc_col] is not None else "",
        }
    return mapping


def extract_pan_from_password_page(page: Page) -> str:
    """Password page par dikhne wala 'PAN APWPK4340L' type text read karta hai."""
    locator = page.locator("text=/PAN\\s*:?\\s*[A-Z]{5}[0-9]{4}[A-Z]/").first
    locator.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    text = locator.inner_text()
    match = re.search(r"[A-Z]{5}[0-9]{4}[A-Z]", text)
    if not match:
        raise RuntimeError(f"Password page se PAN nahi nikal paya. Text mila: {text!r}")
    return match.group(0)


def debug_dump(page: Page, label: str):
    """Kisi bhi step par error aaye to screenshot + page HTML save karta hai,
    taaki exact selector/wording dekh kar fix kiya ja sake."""
    folder = Path(__file__).parent / "debug"
    folder.mkdir(exist_ok=True)
    ts = time.strftime("%Y%m%d_%H%M%S")
    png_path = folder / f"{label}_{ts}.png"
    html_path = folder / f"{label}_{ts}.html"
    try:
        page.screenshot(path=str(png_path), full_page=True)
        print(f"[DEBUG] Screenshot save ki: {png_path}")
    except Exception as e:
        print(f"[DEBUG] Screenshot save nahi ho payi: {e}")
    try:
        html_path.write_text(page.content(), encoding="utf-8")
        print(f"[DEBUG] Page HTML save ki: {html_path}")
    except Exception as e:
        print(f"[DEBUG] HTML save nahi ho payi: {e}")


def dismiss_lingering_overlay(page: Page):
    """Koi purana khula dropdown/overlay ho to Escape + backdrop click se band
    karta hai, taaki asli button tak click pahunch sake."""
    try:
        page.keyboard.press("Escape")
        page.wait_for_timeout(300)
    except Exception:
        pass
    backdrop = page.locator(".cdk-overlay-backdrop").first
    try:
        if backdrop.is_visible():
            backdrop.click(force=True, timeout=1000)
            page.wait_for_timeout(300)
    except Exception:
        pass


def click_with_overlay_retry(locator, page: Page, label: str, attempts: int = 3):
    """Click karta hai; agar koi invisible overlay block kare to overlay
    clear karke dobara try karta hai."""
    last_err = None
    for i in range(attempts):
        try:
            locator.scroll_into_view_if_needed(timeout=ACTION_TIMEOUT_MS)
            locator.click(timeout=ACTION_TIMEOUT_MS)
            return
        except Exception as e:
            last_err = e
            print(f"[WARN] '{label}' click attempt {i + 1} fail hua, overlay clear karke retry...")
            dismiss_lingering_overlay(page)
            page.wait_for_timeout(500)
    raise RuntimeError(f"'{label}' par click nahi ho paya: {last_err}")


def click_mat_radio(page: Page, label_text: str):
    """Angular Material mat-radio-button ko select karta hai. Container par
    click kabhi-kabhi ripple/overlay ki wajah se register nahi hota, isliye
    seedha uske andar wale native <label> (jo input se 'for' se linked hai)
    par click karta hai — ye har browser mein reliably radio ko check karta
    hai."""
    radio_button = page.locator("mat-radio-button", has_text=label_text).first
    radio_button.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    label_el = radio_button.locator("label.mdc-label").first
    if label_el.count() == 0:
        label_el = radio_button
    click_with_overlay_retry(label_el, page, f"radio '{label_text}'")

    # Verify it actually got checked; agar nahi, ek aur try (native input par force click).
    native_input = radio_button.locator("input[type='radio']").first
    try:
        native_input.wait_for(state="attached", timeout=3000)
        if not native_input.is_checked():
            native_input.click(force=True, timeout=ACTION_TIMEOUT_MS)
            page.wait_for_timeout(300)
    except Exception:
        pass


def click_mat_checkbox(page: Page, label_text: str):
    """Angular Material mat-checkbox ko uske visible label text se dhoond kar
    tick karta hai — reliability ke liye native <label> par click karta hai."""
    checkbox = page.locator("mat-checkbox", has_text=label_text).first
    checkbox.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    label_el = checkbox.locator("label").first
    if label_el.count() == 0:
        label_el = checkbox
    click_with_overlay_retry(label_el, page, f"checkbox '{label_text}'")


def select_mat_select_by_formcontrol(page: Page, formcontrolname: str, option_text: str):
    """Angular Material mat-select ko uske formcontrolname se dhoond kar kholta
    hai, fir diye gaye option (partial text match) select karta hai."""
    page.bring_to_front()
    trigger = page.locator(f"mat-select[formcontrolname='{formcontrolname}']").first
    trigger.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    trigger.scroll_into_view_if_needed(timeout=ACTION_TIMEOUT_MS)

    option = page.locator("mat-option", has_text=option_text).first

    # Panel khulne tak click try karo (kabhi kabhi pehla click focus hi leta
    # hai, panel nahi kholta — isliye is check ke sath dobara try karte hain).
    for attempt in range(4):
        click_with_overlay_retry(trigger, page, f"dropdown '{formcontrolname}'")
        try:
            option.wait_for(state="visible", timeout=3000)
            break
        except PWTimeoutError:
            print(f"[WARN] '{formcontrolname}' dropdown panel nahi khula (try {attempt + 1}), "
                  f"keyboard se try kar rahe hain...")
            try:
                trigger.press("Enter")
                option.wait_for(state="visible", timeout=3000)
                break
            except PWTimeoutError:
                continue
    else:
        raise RuntimeError(f"'{formcontrolname}' dropdown panel khul hi nahi raha.")

    click_with_overlay_retry(option, page, f"option '{option_text}'")


def handle_dual_login_popup(page: Page, wait_ms: int = 6000):
    """'Dual Login Detected' / duplicate session popup aaye to 'Login Here'
    button khud dabata hai. Popup na aaye to chup-chaap aage badh jata hai."""
    try:
        login_here_btn = page.get_by_role(
            "button", name=re.compile(r"login\s*here", re.I)
        ).first
        login_here_btn.wait_for(state="visible", timeout=wait_ms)
        login_here_btn.click()
        print("[INFO] 'Dual Login Detected' popup mila — 'Login Here' dabaya.")
        page.wait_for_timeout(500)
    except PWTimeoutError:
        pass  # popup nahi aaya, normal flow jari rakho


def login_flow(page: Page, partner_map: dict):
    print(f"[INFO] Login page khol rahe hain: {LOGIN_URL}")
    page.goto(LOGIN_URL, timeout=ACTION_TIMEOUT_MS)

    print("[ACTION REQUIRED] Apna PAN/User-ID daalkar 'Continue' dabayein.")
    print("[INFO] Password page ka wait kar rahe hain (max 5 min)...")
    page.wait_for_selector("text=Enter password for your e-Filing account", timeout=NAV_TIMEOUT_MS)
    print("[INFO] Password page pe pahunch gaye.")

    pan = extract_pan_from_password_page(page)
    print(f"[INFO] Password page par PAN mila: {pan}")

    record = partner_map.get(pan)
    if not record or not record["password"]:
        raise RuntimeError(f"PAN {pan} ke liye Bank.xlsx mein Password nahi mila.")

    click_mat_checkbox(page, "Please confirm your secure access message displayed above")
    print("[INFO] Secure access message checkbox tick ki.")

    password_field = page.locator("input[type='password']").first
    password_field.fill(record["password"])
    print("[INFO] Password bhar diya.")

    continue_btn = page.get_by_role("button", name="Continue", exact=True).first
    click_with_overlay_retry(continue_btn, page, "Continue (login)")
    print("[INFO] Continue click kiya...")

    # 'Request is not authenticated' jaisi transient error portal kai baar
    # deta hai — jab tak ye error dikhe (ya dashboard na aa jaye), Continue
    # baar baar dabate raho.
    for attempt in range(8):
        try:
            page.wait_for_selector("text=Request is not authenticated", timeout=4000)
        except PWTimeoutError:
            break  # error nahi dikha, matlab aage badh gaye
        print(f"[WARN] 'Request is not authenticated' dikha (try {attempt + 1}) — "
              f"dobara Continue dabate hain...")
        page.wait_for_timeout(800)
        click_with_overlay_retry(continue_btn, page, "Continue (login retry)")

    handle_dual_login_popup(page)

    print("[INFO] Dashboard ka wait kar rahe hain...")
    page.wait_for_url(re.compile(r".*dashboard.*"), timeout=NAV_TIMEOUT_MS)
    page.wait_for_load_state("networkidle")
    print("[INFO] Login safal — Dashboard khul gaya.")
    return pan, record


def open_my_profile(page: Page):
    """Top-right profile dropdown khol kar 'My Profile' click karta hai."""
    profile_dropdown = page.locator("header, .mat-toolbar").get_by_text(re.compile(r"^individual$", re.I)).first
    # Zyada reliable: naam ke upar chevron/dropdown ko user-role text ke paas se dhoondo.
    try:
        profile_dropdown.wait_for(state="visible", timeout=5000)
        # Parent clickable trigger (naam + chevron) usually ek level upar hota hai.
        trigger = profile_dropdown.locator("xpath=ancestor::*[self::button or @role='button' or contains(@class,'clickable')][1]")
        if trigger.count() == 0:
            trigger = profile_dropdown.locator("xpath=..")
        click_with_overlay_retry(trigger.first, page, "Profile dropdown")
    except PWTimeoutError:
        # Fallback: seedha 'My Profile' text dhoondo (kabhi kabhi menu already open hota hai)
        pass

    my_profile_item = page.get_by_text("My Profile", exact=True).first
    try:
        my_profile_item.wait_for(state="visible", timeout=4000)
    except PWTimeoutError:
        # Dropdown khula nahi — naam text par direct click try karo.
        name_locator = page.locator("text=/Individual/i").first
        name_locator.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
        click_with_overlay_retry(name_locator, page, "Profile name (fallback)")
        my_profile_item = page.get_by_text("My Profile", exact=True).first
        my_profile_item.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)

    click_with_overlay_retry(my_profile_item, page, "'My Profile' menu item")
    print("[INFO] 'My Profile' par click kiya.")

    page.wait_for_selector("text=Personal Details", timeout=ACTION_TIMEOUT_MS)
    page.wait_for_load_state("networkidle")
    print("[INFO] Personal Details page khul gaya.")


def open_add_bank_account(page: Page):
    bank_account_link = page.get_by_text("My Bank Account", exact=True).first
    bank_account_link.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    click_with_overlay_retry(bank_account_link, page, "'My Bank Account' sidebar link")
    print("[INFO] 'My Bank Account' par click kiya.")

    page.wait_for_selector("text=My Bank Accounts", timeout=ACTION_TIMEOUT_MS)
    page.wait_for_load_state("networkidle")

    add_btn = page.get_by_role("button", name=re.compile(r"add bank account", re.I)).first
    add_btn.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    click_with_overlay_retry(add_btn, page, "'+ Add Bank Account' button")
    print("[INFO] 'Add Bank Account' form khola.")

    page.wait_for_selector("text=Bank Details", timeout=ACTION_TIMEOUT_MS)


def fill_bank_account_form(page: Page, record: dict):
    bank_account = record["bank_account"]
    ifsc = record["ifsc"]

    page.bring_to_front()
    acc_field = page.locator("input[formcontrolname='bankAcctNum']").first
    acc_field.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    acc_field.click()
    acc_field.fill(bank_account)
    print(f"[INFO] Bank Account Number bhara: {bank_account}")

    confirm_field = page.locator("input[formcontrolname='confirmBankAccountNumber']").first
    confirm_field.click()
    confirm_field.fill(bank_account)
    confirm_field.press("Tab")  # blur trigger karne ke liye — match-validation isi par depend karti hai
    print("[INFO] Confirm Bank Account Number bhara.")

    try:
        page.wait_for_selector("text=Bank Account number and Confirm bank account number is same", timeout=8000)
        print("[INFO] Account number match confirm hua.")
    except PWTimeoutError:
        print("[WARN] Match-confirmation message nahi dikha — aage badh rahe hain "
              "(agar form invalid hoga to 'Proceed' button disabled hi rahega).")

    select_mat_select_by_formcontrol(page, "accTypeCd", "Saving")
    print(f"[INFO] Account Type select kiya: {ACCOUNT_TYPE_LABEL}")

    page.bring_to_front()
    click_mat_radio(page, ACCOUNT_HOLDER_TYPE)
    print(f"[INFO] Account Holder Type select kiya: {ACCOUNT_HOLDER_TYPE}")

    ifsc_field = page.locator("input[formcontrolname='ifscCd']").first
    ifsc_field.click()
    ifsc_field.fill(ifsc)
    ifsc_field.press("Tab")
    print(f"[INFO] IFSC bhara: {ifsc}")

    # Bank Name / Branch auto-populate hone ka wait (portal IFSC lookup karta hai).
    page.wait_for_selector("text=Bank Name", timeout=ACTION_TIMEOUT_MS)
    page.wait_for_timeout(1000)
    try:
        page.wait_for_selector("text=IFSC entered is not valid", timeout=1500)
        raise RuntimeError(f"IFSC '{ifsc}' invalid bataya gaya portal ne — Excel mein IFSC check karein.")
    except PWTimeoutError:
        pass
    print("[INFO] Bank Name/Branch auto-populate ho gaya.")

    click_mat_radio(page, "Yes")
    print("[INFO] 'Nominate this bank account for Refund' = Yes select kiya.")

    page.bring_to_front()
    proceed_btn = page.locator("#addBankAccount").first
    proceed_btn.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)

    # Button tab tak disabled rehta hai jab tak form valid nahi hota — enable
    # hone ka wait karo (max ~15s), fir click karo.
    for _ in range(30):
        if proceed_btn.is_enabled():
            break
        page.wait_for_timeout(500)
    else:
        print("[WARN] 'Proceed to e-Verify' button 15 sec baad bhi disabled hai — "
              "shayad koi field abhi bhi invalid hai. Phir bhi click try kar rahe hain.")

    click_with_overlay_retry(proceed_btn, page, "'Proceed to e-Verify' button")
    print("[INFO] 'Proceed to e-Verify' par click kiya.")


def run_evc_otp_flow(page: Page):
    # 'How do you want to e-verify?' text page par DO baar hai — ek hidden
    # <legend> (accessibility ke liye) aur ek visible heading. Hidden wale ko
    # pehle wait karne se timeout hota hai, isliye seedha radio button ka
    # wait karte hain jo actually visible/interactive hai.
    otp_radio = page.locator(
        "mat-radio-button", has_text="I would like to verify using OTP on mobile number registered with Aadhaar"
    ).first
    otp_radio.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    print("[INFO] e-Verify page khul gaya.")

    click_mat_radio(page, "I would like to verify using OTP on mobile number registered with Aadhaar")
    print("[INFO] 'OTP on mobile number registered with Aadhaar' radio select kiya.")

    page.bring_to_front()
    continue_btn = page.get_by_role("button", name="Continue", exact=True).first
    click_with_overlay_retry(continue_btn, page, "Continue (e-verify method)")
    print("[INFO] Continue click kiya — Aadhaar OTP consent modal ka wait...")

    agree_checkbox = page.locator("mat-checkbox", has_text="I agree to validate my Aadhaar Details").first
    agree_checkbox.wait_for(state="visible", timeout=ACTION_TIMEOUT_MS)
    click_mat_checkbox(page, "I agree to validate my Aadhaar Details")
    print("[INFO] 'I agree to validate my Aadhaar Details' checkbox tick ki.")

    generate_otp_btn = page.get_by_role("button", name=re.compile(r"generate aadhaar otp", re.I)).first
    click_with_overlay_retry(generate_otp_btn, page, "'Generate Aadhaar OTP' button")
    print("[INFO] 'Generate Aadhaar OTP' click kar diya.")

    print("\n[ACTION REQUIRED] OTP aapke Aadhaar-registered mobile number par aa gaya hoga.")
    print("[ACTION REQUIRED] Browser mein OTP manually daalein aur 'Validate' dabayein.")


def _automate(page: Page, partner_map: dict):
    pan, record = login_flow(page, partner_map)
    page.bring_to_front()
    open_my_profile(page)
    page.bring_to_front()
    open_add_bank_account(page)
    page.bring_to_front()
    fill_bank_account_form(page, record)
    page.bring_to_front()
    run_evc_otp_flow(page)


def run():
    partner_map = load_partner_map(EXCEL_PATH)
    print(f"[INFO] Excel se {len(partner_map)} PAN entries load hui.")

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=False,
            channel="chrome",
            args=["--disable-blink-features=AutomationControlled"],
        )
        context = browser.new_context(no_viewport=True)
        context.add_init_script(
            "Object.defineProperty(navigator, 'webdriver', {get: () => undefined});"
        )
        page = context.new_page()
        page.set_default_timeout(ACTION_TIMEOUT_MS)

        try:
            _automate(page, partner_map)
        except Exception as e:
            print(f"\n[ERROR] Automation mein error aayi: {e}")
            debug_dump(page, "error")
            print("[INFO] Upar wali 'debug' folder mein screenshot + HTML dekh kar mujhe bhej dena.")

        input("\nBrowser abhi open rahega — manually OTP daal kar Validate dabayein. "
              "Poora ho jaye to yahan Enter dabayein (browser band ho jayega)... ")

        browser.close()
        print("[INFO] Browser band kar diya. Script complete.")


if __name__ == "__main__":
    run()